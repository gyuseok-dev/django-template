"""Helper functions for creating mock test data"""

from io import BytesIO

import openpyxl
from django.core.files.uploadedfile import InMemoryUploadedFile


def create_mock_medical_record_excel(rows=2):
    """
    환자 진료 기록 mock 엑셀 파일 생성

    Args:
        rows: 생성할 데이터 행 수 (기본값: 2)

    Returns:
        InMemoryUploadedFile: 테스트용 엑셀 파일
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # 헤더 행
    ws.append(
        [
            "챠트번호",
            "수진자명",
            "주민번호",
            "보험유형",
            "진료과",
            "담당의",
            "초/재",
            "내원일",
        ]
    )

    # 데이터 행 생성
    for i in range(rows):
        ws.append(
            [
                10001 + i,  # 챠트번호
                f"환자{i}",  # 수진자명
                "8001011234567",  # 주민번호
                "국민건강보험",  # 보험유형
                "내과",  # 진료과
                "김의사",  # 담당의
                "신환" if i % 2 == 0 else "재진",  # 초/재
                "202401151030",  # 내원일 (YYYYMMDDHHmm)
            ]
        )

    # BytesIO로 변환
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    return InMemoryUploadedFile(
        excel_file,
        "medical_record_file",
        "test_medical.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        excel_file.getbuffer().nbytes,
        None,
    )


def create_mock_visit_channel_excel(rows=2):
    """
    내원 경로 mock 엑셀 파일 생성

    Args:
        rows: 생성할 데이터 행 수 (기본값: 2)

    Returns:
        InMemoryUploadedFile: 테스트용 엑셀 파일
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # 헤더 행
    ws.append(
        [
            "차트번호",
            "이름",
            "보험유형",
            "진료과목",
            "담당의",
            "초재진구분",
            "내원경로",
            "내원날짜",
        ]
    )

    # 데이터 행 생성
    channels = ["인터넷", "간판", "소개", "전화", "재방문", "기타"]
    for i in range(rows):
        ws.append(
            [
                10001 + i,  # 차트번호
                f"환자{i}",  # 이름
                "국민건강보험",  # 보험유형
                "내과",  # 진료과목
                "김의사",  # 담당의
                "신환" if i % 2 == 0 else "재진",  # 초재진구분
                channels[i % len(channels)],  # 내원경로
                "202401151030",  # 내원날짜 (YYYYMMDDHHmm)
            ]
        )

    # 통계 행 추가 (제거되어야 하는 행)
    ws.append(["【 소 계 】", "", "", "", "", "", "", ""])
    ws.append(["【 총 계 】", "", "", "", "", "", "", ""])

    # BytesIO로 변환
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    return InMemoryUploadedFile(
        excel_file,
        "visit_source_file",
        "test_visit.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        excel_file.getbuffer().nbytes,
        None,
    )


def create_invalid_excel():
    """
    잘못된 형식의 mock 엑셀 파일 생성 (에러 테스트용)

    Returns:
        InMemoryUploadedFile: 잘못된 형식의 테스트용 엑셀 파일
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # 잘못된 헤더
    ws.append(["잘못된", "헤더", "형식"])

    # 데이터가 맞지 않는 행
    ws.append(["데이터1", "데이터2"])

    # BytesIO로 변환
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    return InMemoryUploadedFile(
        excel_file,
        "medical_record_file",
        "test_invalid.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        excel_file.getbuffer().nbytes,
        None,
    )


def create_mock_medical_with_empty_rows():
    """
    빈 행이 포함된 환자 진료 기록 엑셀 생성 (빈 행 제거 테스트용)

    Returns:
        InMemoryUploadedFile: 빈 행이 포함된 테스트용 엑셀 파일
    """
    wb = openpyxl.Workbook()
    ws = wb.active

    # 헤더
    ws.append(
        [
            "챠트번호",
            "수진자명",
            "주민번호",
            "보험유형",
            "진료과",
            "담당의",
            "초/재",
            "내원일",
        ]
    )

    # 데이터 행
    ws.append(
        [
            10001,
            "환자1",
            "8001011234567",
            "국민건강보험",
            "내과",
            "김의사",
            "신환",
            "202401151030",
        ]
    )

    # 빈 행
    ws.append([None, None, None, None, None, None, None, None])
    ws.append(["", "", "", "", "", "", "", ""])

    # 데이터 행
    ws.append(
        [
            10002,
            "환자2",
            "9001011234567",
            "보호1종",
            "외과",
            "이의사",
            "재진",
            "202401151100",
        ]
    )

    # BytesIO로 변환
    excel_file = BytesIO()
    wb.save(excel_file)
    excel_file.seek(0)

    return InMemoryUploadedFile(
        excel_file,
        "medical_record_file",
        "test_with_empty.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        excel_file.getbuffer().nbytes,
        None,
    )
